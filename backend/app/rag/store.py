"""Document loading, chunking, and FAISS index management."""

from __future__ import annotations

import json
import shutil
import uuid
from dataclasses import asdict, dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from langchain_community.docstore.in_memory import InMemoryDocstore
from langchain_community.vectorstores import FAISS
from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter
from pypdf import PdfReader

from app.config import Settings, get_settings
from app.rag.embeddings import get_embeddings

try:
    import faiss  # type: ignore
except ImportError:  # pragma: no cover
    faiss = None


@dataclass
class DocumentMeta:
    id: str
    filename: str
    title: str
    char_count: int
    chunk_count: int
    uploaded_at: str
    source_path: str


class KnowledgeStore:
    def __init__(self, settings: Settings | None = None) -> None:
        self.settings = settings or get_settings()
        self.meta_path = self.settings.data_dir / "documents_meta.json"
        self._vectorstore: FAISS | None = None
        self._documents: dict[str, DocumentMeta] = {}
        self._load_meta()
        self._load_or_create_index()

    def _load_meta(self) -> None:
        if self.meta_path.exists():
            raw = json.loads(self.meta_path.read_text(encoding="utf-8"))
            self._documents = {item["id"]: DocumentMeta(**item) for item in raw}
        else:
            self._documents = {}

    def _save_meta(self) -> None:
        payload = [asdict(doc) for doc in self._documents.values()]
        self.meta_path.write_text(json.dumps(payload, indent=2), encoding="utf-8")

    def _embedding_dim(self) -> int:
        probe = get_embeddings().embed_query("dimension probe")
        return len(probe)

    def _create_empty_index(self) -> FAISS:
        embeddings = get_embeddings()
        dim = self._embedding_dim()
        if faiss is None:
            raise RuntimeError("faiss is required")
        index = faiss.IndexFlatL2(dim)
        return FAISS(
            embedding_function=embeddings,
            index=index,
            docstore=InMemoryDocstore({}),
            index_to_docstore_id={},
        )

    def _load_or_create_index(self) -> None:
        index_file = self.settings.index_dir / "index.faiss"
        if index_file.exists():
            try:
                store = FAISS.load_local(
                    str(self.settings.index_dir),
                    get_embeddings(),
                    allow_dangerous_deserialization=True,
                )
                # Rebuild if embedding provider/dimension changed (e.g. mock → TAMU).
                expected = self._embedding_dim()
                if store.index.d != expected:
                    raise ValueError(f"index dim {store.index.d} != embedding dim {expected}")
                self._vectorstore = store
                return
            except Exception:
                shutil.rmtree(self.settings.index_dir, ignore_errors=True)
                self.settings.index_dir.mkdir(parents=True, exist_ok=True)
                # Drop stale meta so samples re-seed with the new embeddings.
                if self.meta_path.exists():
                    self.meta_path.unlink()
                self._documents = {}
                # Also clear copied docs so seed starts clean.
                for path in self.settings.docs_dir.glob("*"):
                    if path.is_file():
                        path.unlink()
        self._vectorstore = self._create_empty_index()
        self._persist()

    def _persist(self) -> None:
        assert self._vectorstore is not None
        self._vectorstore.save_local(str(self.settings.index_dir))
        self._save_meta()

    @property
    def vectorstore(self) -> FAISS:
        assert self._vectorstore is not None
        return self._vectorstore

    def list_documents(self) -> list[DocumentMeta]:
        return sorted(self._documents.values(), key=lambda d: d.uploaded_at, reverse=True)

    def document_count(self) -> int:
        return len(self._documents)

    def chunk_count(self) -> int:
        return sum(d.chunk_count for d in self._documents.values())

    def _read_file(self, path: Path) -> str:
        suffix = path.suffix.lower()
        if suffix in {".txt", ".md", ".markdown", ".csv"}:
            return path.read_text(encoding="utf-8", errors="ignore")
        if suffix == ".pdf":
            reader = PdfReader(str(path))
            return "\n".join(page.extract_text() or "" for page in reader.pages)
        if suffix in {".docx"}:
            from docx import Document as DocxDocument

            doc = DocxDocument(str(path))
            return "\n".join(p.text for p in doc.paragraphs)
        raise ValueError(f"Unsupported file type: {suffix}")

    def ingest_file(self, source: Path, original_name: str | None = None) -> DocumentMeta:
        filename = original_name or source.name
        text = self._read_file(source).strip()
        if not text:
            raise ValueError("Document is empty after extraction")

        doc_id = str(uuid.uuid4())
        dest = self.settings.docs_dir / f"{doc_id}_{filename}"
        shutil.copy2(source, dest)

        splitter = RecursiveCharacterTextSplitter(
            chunk_size=self.settings.chunk_size,
            chunk_overlap=self.settings.chunk_overlap,
            separators=["\n\n", "\n", ". ", " ", ""],
        )
        chunks = splitter.split_text(text)
        documents = [
            Document(
                page_content=chunk,
                metadata={
                    "doc_id": doc_id,
                    "filename": filename,
                    "chunk_index": idx,
                    "source": filename,
                },
            )
            for idx, chunk in enumerate(chunks)
        ]
        self.vectorstore.add_documents(documents)

        meta = DocumentMeta(
            id=doc_id,
            filename=filename,
            title=Path(filename).stem.replace("_", " ").title(),
            char_count=len(text),
            chunk_count=len(documents),
            uploaded_at=datetime.now(timezone.utc).isoformat(),
            source_path=str(dest),
        )
        self._documents[doc_id] = meta
        self._persist()
        return meta

    def ingest_text(self, title: str, text: str) -> DocumentMeta:
        safe_name = f"{title.replace(' ', '_').lower()}.txt"
        temp = self.settings.docs_dir / f"_tmp_{uuid.uuid4().hex}.txt"
        temp.write_text(text, encoding="utf-8")
        try:
            return self.ingest_file(temp, original_name=safe_name)
        finally:
            if temp.exists():
                temp.unlink()

    def delete_document(self, doc_id: str) -> bool:
        if doc_id not in self._documents:
            return False
        meta = self._documents.pop(doc_id)
        path = Path(meta.source_path)
        if path.exists():
            path.unlink()

        remaining = list(self._documents.values())
        self._documents = {}
        self._vectorstore = self._create_empty_index()
        for item in remaining:
            src = Path(item.source_path)
            if not src.exists():
                continue
            text = self._read_file(src).strip()
            if not text:
                continue
            splitter = RecursiveCharacterTextSplitter(
                chunk_size=self.settings.chunk_size,
                chunk_overlap=self.settings.chunk_overlap,
                separators=["\n\n", "\n", ". ", " ", ""],
            )
            chunks = splitter.split_text(text)
            documents = [
                Document(
                    page_content=chunk,
                    metadata={
                        "doc_id": item.id,
                        "filename": item.filename,
                        "chunk_index": idx,
                        "source": item.filename,
                    },
                )
                for idx, chunk in enumerate(chunks)
            ]
            self.vectorstore.add_documents(documents)
            item.chunk_count = len(documents)
            item.char_count = len(text)
            self._documents[item.id] = item
        self._persist()
        return True

    def similarity_search(self, query: str, k: int | None = None) -> list[Document]:
        k = k or self.settings.top_k
        if self.chunk_count() == 0:
            return []
        return self.vectorstore.similarity_search(query, k=min(k, max(self.chunk_count(), 1)))

    def similarity_search_with_score(self, query: str, k: int | None = None) -> list[tuple[Document, float]]:
        k = k or self.settings.top_k
        if self.chunk_count() == 0:
            return []
        return self.vectorstore.similarity_search_with_score(query, k=min(k, max(self.chunk_count(), 1)))

    def seed_samples_if_empty(self) -> list[DocumentMeta]:
        if self.document_count() > 0:
            return []
        sample_dir = self.settings.sample_docs_dir
        added: list[DocumentMeta] = []
        if not sample_dir.exists():
            return added
        for path in sorted(sample_dir.glob("*")):
            if path.is_file() and path.suffix.lower() in {".txt", ".md", ".pdf", ".docx"}:
                added.append(self.ingest_file(path))
        return added


_store: KnowledgeStore | None = None


def get_store() -> KnowledgeStore:
    global _store
    if _store is None:
        _store = KnowledgeStore()
        _store.seed_samples_if_empty()
    return _store


def reset_store() -> KnowledgeStore:
    global _store
    _store = KnowledgeStore()
    return _store
